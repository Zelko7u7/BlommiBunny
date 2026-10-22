"""Validate the original workbook and import an inventory snapshot, without replaying sales.

Run: python tools/import_inventory.py [--apply]
Outputs and the pre-import database backup stay under .local/inventory-import (not web assets).
Existing products are never overwritten; conflicting data aborts the whole transaction.
"""
import argparse
from collections import defaultdict
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import unicodedata
from datetime import datetime
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
ALIASES = {
    'Sobre Mascarilla Stickers Para Acné Happy Clean': 'Stickers Para Acné Happy Clean',
    'Tinta Gloss Corazon Trendy': 'Tinta Gloss Corazón',
    'Rubor Doble Mia': 'Mia Blush, Rubor doble',
    'Pompon Cotton': 'Pompon Iluminador Cotton',
    'Serum Hidratante Morado Trendy': 'Serum Hidratante Ácido Hialurónico',
}


def normalize(value):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFD', str(value)).encode('ascii', 'ignore').decode().lower())


def number(value, location, blank_zero=False, integer=False):
    if value is None and blank_zero:
        return Decimal(0)
    if value is None or isinstance(value, bool):
        raise ValueError(f'{location}: missing or invalid number: {value!r}')
    try:
        result = Decimal(str(value))
    except InvalidOperation as error:
        raise ValueError(f'{location}: invalid number {value!r}') from error
    if not result.is_finite() or (integer and result != result.to_integral_value()):
        raise ValueError(f'{location}: number must be finite' + (' and integral' if integer else ''))
    return result


def prepare(workbook, image_manifest):
    values = load_workbook(workbook, data_only=True)
    sheet = values['Productos']
    expected = ['Código', 'Producto', 'Categoría', 'Costo', 'Precio', 'Stock inicial', 'Compras', 'Ventas', 'Ajuste manual', 'Stock actual', 'Stock mínimo', 'Estado', 'Ganancia/u', 'Valor stock']
    if [sheet.cell(4, c).value for c in range(1, 15)] != expected:
        raise ValueError('Unexpected Productos columns. Review the workbook before importing.')
    movements = {}
    for title in ['Compras', 'Ventas']:
        totals = defaultdict(Decimal)
        for row in values[title].iter_rows(min_row=5):
            if row[3].value:
                quantity = number(row[4].value, f'{title}!E{row[4].row}', integer=True)
                if quantity < 0:
                    raise ValueError(f'{title}: negative quantity')
                totals[normalize(row[3].value)] += quantity
        movements[title] = totals
    image_by_name = {normalize(image['name']): image for image in image_manifest}
    products, seen_names, seen_codes, matched, warnings = [], set(), set(), set(), []
    for r in range(5, sheet.max_row + 1):
        row = [sheet.cell(r, c).value for c in range(1, 15)]
        if not row[1]:
            if any(row[c] not in (None, '') for c in [0, 2, 3, 4, 5, 8, 10]):
                raise ValueError(f'Productos row {r}: data without a product name')
            continue
        name, category = str(row[1]).strip(), str(row[2] or '').strip()
        key = normalize(name)
        code = str(row[0]).strip() if row[0] is not None else 'BB-XLS-' + hashlib.sha256(key.encode()).hexdigest()[:12].upper()
        if not key or key in seen_names or code in seen_codes:
            raise ValueError(f'Productos row {r}: empty or duplicate name/code')
        if not category or len(name) > 140 or len(category) > 120 or len(code) > 60:
            raise ValueError(f'Productos row {r}: missing category or field too long')
        seen_names.add(key); seen_codes.add(code)
        cost = number(row[3], f'D{r}')
        price = number(row[4], f'E{r}')
        initial = number(row[5], f'F{r}', blank_zero=True, integer=True)
        purchases = number(row[6], f'G{r}', integer=True)
        sales = number(row[7], f'H{r}', integer=True)
        adjustment = number(row[8], f'I{r}', blank_zero=True, integer=True)
        stock = number(row[9], f'J{r}', integer=True)
        minimum = number(row[10], f'K{r}', blank_zero=True, integer=True)
        gain = number(row[12], f'M{r}')
        stock_value = number(row[13], f'N{r}')
        if purchases != movements['Compras'][key] or sales != movements['Ventas'][key]:
            raise ValueError(f'{name}: cached purchases/sales do not agree with source movements')
        if stock != initial + purchases - sales + adjustment or gain != price - cost or stock_value != stock * cost:
            raise ValueError(f'{name}: cached stock, profit or inventory valuation is inconsistent')
        if min(cost, price, stock, minimum) < 0 or price <= 0:
            raise ValueError(f'{name}: negative values or nonpositive sale price')
        if any(v != v.quantize(Decimal('.01')) or v >= Decimal('1000000000000') for v in (cost, price)):
            raise ValueError(f'{name}: amount outside database precision')
        if stock > 2147483647 or minimum > 2147483647:
            raise ValueError(f'{name}: quantity exceeds database range')
        photo_name = ALIASES.get(name, name)
        image = image_by_name.get(normalize(photo_name))
        if image:
            if image['imageUrl'] in matched:
                raise ValueError(f'{name}: photograph assigned to more than one product')
            matched.add(image['imageUrl'])
        if row[10] is None:
            warnings.append(f'{name}: stock mínimo vacío; se guarda 0 (sin umbral de stock bajo).')
        products.append(dict(sourceRow=r, name=name, code=code, category=category, purchasePrice=str(cost), salePrice=str(price), stock=int(stock), minimumStock=int(minimum), imageUrl=image['imageUrl'] if image else None, unitProfit=str(gain), stockValue=str(stock_value)))
    unknown = (set(movements['Compras']) | set(movements['Ventas'])) - seen_names
    if unknown:
        raise ValueError(f'Movements reference products missing from Productos: {sorted(unknown)}')
    return dict(source=workbook.name, sourceSha256=hashlib.sha256(workbook.read_bytes()).hexdigest(), products=products, summary=dict(products=len(products), units=sum(p['stock'] for p in products), inventoryCostValue=str(sum(Decimal(p['stockValue']) for p in products)), availableProducts=sum(p['stock'] > 0 for p in products), outOfStock=sum(p['stock'] == 0 for p in products), matchedPhotos=len(matched)), unassignedPhotos=[image['name'] for image in image_manifest if image['imageUrl'] not in matched], warnings=warnings)


def quote(value):
    return 'NULL' if value is None else "'" + str(value).replace("'", "''") + "'"


def build_sql(report):
    rows = ["(" + ','.join([quote(p['name']), quote(p['code']), quote(p['category']), p['purchasePrice'], p['salePrice'], str(p['stock']), str(p['minimumStock']), quote(p['imageUrl'])]) + ")" for p in report['products']]
    reference = quote('Inventario inicial Excel SHA256 ' + report['sourceSha256'])
    return """BEGIN;
SELECT pg_advisory_xact_lock(73194927);
CREATE TEMP TABLE source_inventory (name text,code text,category text,cost numeric,price numeric,stock integer,minimum_stock integer,image_url text) ON COMMIT DROP;
INSERT INTO source_inventory VALUES
""" + ',\n'.join(rows) + """;
-- Never overwrite a product changed since the initial import.
DO $$ BEGIN
 IF EXISTS (SELECT 1 FROM source_inventory s JOIN product p ON p.code=s.code OR lower(trim(p.name))=lower(s.name)
  JOIN category c ON c.id=p.category_id
  WHERE p.name<>s.name OR p.code<>s.code OR c.name<>s.category OR p.purchase_price<>s.cost OR p.sale_price<>s.price OR p.stock<>s.stock OR p.minimum_stock<>s.minimum_stock OR p.image_url IS DISTINCT FROM s.image_url)
 THEN RAISE EXCEPTION 'Existing inventory differs from Excel; no data was overwritten'; END IF;
END $$;
INSERT INTO category(name,active,created_at,updated_at)
SELECT DISTINCT category,true,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP FROM source_inventory
ON CONFLICT(name) DO NOTHING;
WITH inserted AS (
 INSERT INTO product(name,code,category_id,purchase_price,sale_price,stock,minimum_stock,image_url,active,created_at,updated_at)
 SELECT s.name,s.code,c.id,s.cost,s.price,s.stock,s.minimum_stock,s.image_url,true,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP
 FROM source_inventory s JOIN category c ON c.name=s.category
 WHERE NOT EXISTS (SELECT 1 FROM product p WHERE p.code=s.code OR lower(trim(p.name))=lower(s.name))
 RETURNING id,stock
)
INSERT INTO inventory_movement(product_id,type,quantity,resulting_stock,reference,created_at,updated_at)
SELECT id,'ADJUSTMENT_IN',stock,stock,""" + reference + """,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP FROM inserted;
SELECT count(*) AS imported_products,sum(p.stock) AS units,sum(p.stock*p.purchase_price) AS inventory_value FROM product p JOIN source_inventory s ON p.code=s.code;
COMMIT;
"""


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook', type=Path)
    parser.add_argument('--apply', action='store_true')
    args = parser.parse_args()
    sources = list(ROOT.glob('*.xlsx'))
    if args.workbook is None and len(sources) != 1:
        raise ValueError('Specify --workbook when there is not exactly one Excel file')
    workbook = args.workbook or sources[0]
    manifest = json.loads((ROOT / 'BloomiFront/src/data/product-images.json').read_text(encoding='utf-8'))
    report = prepare(workbook, manifest)
    if not report['products']:
        raise ValueError('Empty inventory; nothing to import')
    output = ROOT / '.local/inventory-import'
    output.mkdir(parents=True, exist_ok=True)
    (output / 'preview.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    sql_file = output / 'import.sql'
    sql_file.write_text(build_sql(report), encoding='utf-8')
    print(json.dumps(report['summary'], ensure_ascii=False))
    print('Fotos pendientes:', ', '.join(report['unassignedPhotos']))
    if args.apply:
        settings = {}
        for line in (ROOT / 'BloomiBack/src/main/resources/application.properties').read_text(encoding='utf-8').splitlines():
            if line and not line.startswith('#') and '=' in line:
                key, value = line.split('=', 1); settings[key] = value
        match = re.fullmatch(r'jdbc:postgresql://([^:/]+):(\d+)/([^/?]+)', settings['spring.datasource.url'])
        if not match:
            raise ValueError('Unsupported database URL')
        host, port, database = match.groups()
        common = ['-h', host, '-p', port, '-U', settings['spring.datasource.username'], '-d', database]
        env = {**os.environ, 'PGPASSWORD': settings['spring.datasource.password'], 'PGCLIENTENCODING': 'UTF8'}
        pg_bin = Path('C:/Program Files/PostgreSQL/18/bin')
        backup = output / ('before-import-' + datetime.now().strftime('%Y%m%d-%H%M%S') + '.dump')
        subprocess.run([str(pg_bin / 'pg_dump.exe'), *common, '-Fc', '-f', str(backup)], env=env, check=True)
        subprocess.run([str(pg_bin / 'psql.exe'), *common, '-v', 'ON_ERROR_STOP=1', '-f', str(sql_file)], env=env, check=True)
        print('Importación terminada. Respaldo:', backup.relative_to(ROOT))


if __name__ == '__main__':
    main()
